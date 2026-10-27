import io

import openpyxl
from docx import Document
from fastapi import status
from fastapi.testclient import TestClient

from app.services.rota_export import sanitize_cell


def _setup_user_and_workspace(client: TestClient, email: str, name: str) -> tuple[dict[str, str], str]:
    client.post(
        "/api/v1/auth/register",
        json={"email": email, "password": "StrongPassword123!", "first_name": name, "last_name": "Manager"},
    )
    login_resp = client.post(
        "/api/v1/auth/login",
        json={"email": email, "password": "StrongPassword123!"},
    )
    token = login_resp.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    ws_resp = client.post(
        "/api/v1/workspaces",
        json={"name": f"{name}'s Workspace"},
        headers=headers,
    )
    ws_id = ws_resp.json()["id"]
    return headers, ws_id


def test_formula_injection_sanitization():
    """Verify formula injection sanitisation covers triggers, whitespace, control characters, and preserves numerics."""
    # Untrusted triggers
    assert sanitize_cell("=1+1") == "'=1+1"
    assert sanitize_cell("+12345") == "'+12345"
    assert sanitize_cell("-calc") == "'-calc"
    assert sanitize_cell("@SUM(A1:A10)") == "'@SUM(A1:A10)"

    # Leading whitespace before triggers
    assert sanitize_cell("   =CMD()") == "'   =CMD()"
    assert sanitize_cell(" \t +123") == "' \t +123"

    # Leading control characters before triggers
    assert sanitize_cell("\t+malicious") == "'\t+malicious"
    assert sanitize_cell("\r@DDE()") == "'\r@DDE()"
    assert sanitize_cell("\n-test") == "'\n-test"

    # Legitimate text
    assert sanitize_cell("Alice Smith") == "Alice Smith"
    assert sanitize_cell("General Skill") == "General Skill"

    # Legitimate numbers (preserved as numbers, not quoted strings)
    assert sanitize_cell(42) == 42
    assert sanitize_cell(7.5) == 7.5
    assert sanitize_cell(0) == 0


def test_export_formats_and_content_verification(client: TestClient):
    """Test CSV, XLSX, PDF, and DOCX exports with both staff and day views, and notes privacy."""
    headers, ws_id = _setup_user_and_workspace(client, "export_test@example.com", "Exporter")

    # Create staff
    staff_resp = client.post(
        f"/api/v1/workspaces/{ws_id}/staff",
        json={"first_name": "Chloé", "last_name": "Müller", "category": "full_time"},
        headers=headers,
    )
    staff_id = staff_resp.json()["id"]

    # Create rota
    rota_resp = client.post(
        f"/api/v1/workspaces/{ws_id}/rotas",
        json={"name": "Export Test Rota", "start_date": "2026-11-02", "end_date": "2026-11-08"},
        headers=headers,
    )
    rota_id = rota_resp.json()["id"]

    # Add a shift with notes and special Unicode characters
    client.post(
        f"/api/v1/workspaces/{ws_id}/rotas/{rota_id}/shifts",
        json={
            "staff_id": staff_id,
            "date": "2026-11-02",
            "start_time": "09:00",
            "end_time": "17:00",
            "notes": "Secret private manager note",
        },
        headers=headers,
    )

    # 1. Test CSV Export (Staff view, notes excluded by default)
    csv_resp = client.get(
        f"/api/v1/workspaces/{ws_id}/rotas/{rota_id}/export?format=csv&view=staff",
        headers=headers,
    )
    assert csv_resp.status_code == status.HTTP_200_OK
    assert "text/csv" in csv_resp.headers["content-type"]
    assert "attachment; filename=" in csv_resp.headers["content-disposition"]
    # Verify UTF-8 BOM
    assert csv_resp.content.startswith(b"\xef\xbb\xbf")
    csv_text = csv_resp.content.decode("utf-8-sig")
    assert "Chloé Müller" in csv_text
    assert "09:00" in csv_text
    assert "17:00" in csv_text
    # Notes must be omitted by default
    assert "Secret private manager note" not in csv_text

    # 2. Test CSV Export with include_notes=true
    csv_notes_resp = client.get(
        f"/api/v1/workspaces/{ws_id}/rotas/{rota_id}/export?format=csv&view=staff&include_notes=true",
        headers=headers,
    )
    assert csv_notes_resp.status_code == status.HTTP_200_OK
    csv_notes_text = csv_notes_resp.content.decode("utf-8-sig")
    assert "Secret private manager note" in csv_notes_text

    # 3. Test CSV Day View
    csv_day_resp = client.get(
        f"/api/v1/workspaces/{ws_id}/rotas/{rota_id}/export?format=csv&view=day",
        headers=headers,
    )
    assert csv_day_resp.status_code == status.HTTP_200_OK
    assert "Time Band" in csv_day_resp.content.decode("utf-8-sig")

    # 4. Test XLSX Export
    xlsx_resp = client.get(
        f"/api/v1/workspaces/{ws_id}/rotas/{rota_id}/export?format=xlsx&view=staff",
        headers=headers,
    )
    assert xlsx_resp.status_code == status.HTTP_200_OK
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in xlsx_resp.headers["content-type"]
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_resp.content))
    ws = wb.active
    assert ws is not None
    # Verify title and staff row
    assert "Export Test Rota" in str(ws["A1"].value)
    found_staff = False
    for row in ws.iter_rows(values_only=True):
        if "Chloé Müller" in row:
            found_staff = True
            break
    assert found_staff

    # 5. Test PDF Export
    pdf_resp = client.get(
        f"/api/v1/workspaces/{ws_id}/rotas/{rota_id}/export?format=pdf&view=staff",
        headers=headers,
    )
    assert pdf_resp.status_code == status.HTTP_200_OK
    assert "application/pdf" in pdf_resp.headers["content-type"]
    assert pdf_resp.content.startswith(b"%PDF-")
    assert len(pdf_resp.content) > 1000

    # 6. Test DOCX Export
    docx_resp = client.get(
        f"/api/v1/workspaces/{ws_id}/rotas/{rota_id}/export?format=docx&view=staff",
        headers=headers,
    )
    assert docx_resp.status_code == status.HTTP_200_OK
    assert "application/vnd.openxmlformats-officedocument.wordprocessingml.document" in docx_resp.headers["content-type"]
    doc = Document(io.BytesIO(docx_resp.content))
    assert len(doc.tables) >= 1
    table_text = " ".join(cell.text for row in doc.tables[0].rows for cell in row.cells)
    assert "Chloé Müller" in table_text

    # 7. Invalid format or view parameter
    bad_fmt = client.get(
        f"/api/v1/workspaces/{ws_id}/rotas/{rota_id}/export?format=unsupported",
        headers=headers,
    )
    assert bad_fmt.status_code == status.HTTP_400_BAD_REQUEST


def test_publishing_lifecycle_and_immutability_enforcement(client: TestClient):
    """Prove that published rotas are strictly immutable across all API mutation endpoints."""
    headers, ws_id = _setup_user_and_workspace(client, "pub_lifecycle@example.com", "Publisher")

    staff_resp = client.post(
        f"/api/v1/workspaces/{ws_id}/staff",
        json={"first_name": "Bob", "last_name": "Builder", "category": "full_time"},
        headers=headers,
    )
    staff_id = staff_resp.json()["id"]

    rota_resp = client.post(
        f"/api/v1/workspaces/{ws_id}/rotas",
        json={"name": "Lifecycle Rota", "start_date": "2026-11-02", "end_date": "2026-11-08"},
        headers=headers,
    )
    rota_id = rota_resp.json()["id"]

    # 1. Attempting to approve an empty rota must fail
    empty_approve = client.post(
        f"/api/v1/workspaces/{ws_id}/rotas/{rota_id}/approve",
        headers=headers,
    )
    assert empty_approve.status_code == status.HTTP_400_BAD_REQUEST
    assert "empty" in empty_approve.json()["detail"].lower()

    # 1b. Regression test: Attempting to bypass approval via general PATCH with status='published' must fail
    empty_patch_bypass = client.patch(
        f"/api/v1/workspaces/{ws_id}/rotas/{rota_id}",
        json={"status": "published"},
        headers=headers,
    )
    assert empty_patch_bypass.status_code == status.HTTP_422_UNPROCESSABLE_ENTITY
    # Verify rota remains in draft and was not published
    check_draft = client.get(
        f"/api/v1/workspaces/{ws_id}/rotas/{rota_id}",
        headers=headers,
    )
    assert check_draft.json()["status"] == "draft"
    assert check_draft.json()["published_at"] is None

    # Add a shift to the rota
    shift_resp = client.post(
        f"/api/v1/workspaces/{ws_id}/rotas/{rota_id}/shifts",
        json={
            "staff_id": staff_id,
            "date": "2026-11-02",
            "start_time": "09:00",
            "end_time": "17:00",
        },
        headers=headers,
    )
    shift_id = shift_resp.json()["shift"]["id"]

    # 2. Approve and publish the rota
    approve_resp = client.post(
        f"/api/v1/workspaces/{ws_id}/rotas/{rota_id}/approve",
        headers=headers,
    )
    assert approve_resp.status_code == status.HTTP_200_OK
    pub_data = approve_resp.json()
    assert pub_data["status"] == "published"
    assert pub_data["published_at"] is not None
    assert pub_data["published_by_user_id"] is not None

    # 3. Prove all mutation endpoints reject changes on the published rota:
    # 3a. Rename via PATCH
    rename_resp = client.patch(
        f"/api/v1/workspaces/{ws_id}/rotas/{rota_id}",
        json={"name": "New Name"},
        headers=headers,
    )
    assert rename_resp.status_code == status.HTTP_400_BAD_REQUEST

    # 3b. Freeform status jumping via PATCH is forbidden
    status_patch = client.patch(
        f"/api/v1/workspaces/{ws_id}/rotas/{rota_id}",
        json={"status": "draft"},
        headers=headers,
    )
    assert status_patch.status_code == status.HTTP_422_UNPROCESSABLE_ENTITY

    # 3c. Update staffing requirements
    req_resp = client.put(
        f"/api/v1/workspaces/{ws_id}/rotas/{rota_id}/requirements",
        json={"requirements": [{"date": "2026-11-02", "bands": []}]},
        headers=headers,
    )
    assert req_resp.status_code == status.HTTP_400_BAD_REQUEST

    # 3d. Copy staffing requirements
    copy_req_resp = client.post(
        f"/api/v1/workspaces/{ws_id}/rotas/{rota_id}/copy-requirements",
        json={"source_date": "2026-11-02", "target_dates": ["2026-11-03"]},
        headers=headers,
    )
    assert copy_req_resp.status_code == status.HTTP_400_BAD_REQUEST

    # 3e. Create shift
    create_shift_resp = client.post(
        f"/api/v1/workspaces/{ws_id}/rotas/{rota_id}/shifts",
        json={
            "staff_id": staff_id,
            "date": "2026-11-03",
            "start_time": "10:00",
            "end_time": "18:00",
        },
        headers=headers,
    )
    assert create_shift_resp.status_code == status.HTTP_400_BAD_REQUEST

    # 3f. Edit shift
    edit_shift_resp = client.patch(
        f"/api/v1/workspaces/{ws_id}/rotas/{rota_id}/shifts/{shift_id}",
        json={"start_time": "10:00"},
        headers=headers,
    )
    assert edit_shift_resp.status_code == status.HTTP_400_BAD_REQUEST

    # 3g. Delete shift
    del_shift_resp = client.delete(
        f"/api/v1/workspaces/{ws_id}/rotas/{rota_id}/shifts/{shift_id}",
        headers=headers,
    )
    assert del_shift_resp.status_code == status.HTTP_400_BAD_REQUEST

    # 3h. Copy shift
    copy_shift_resp = client.post(
        f"/api/v1/workspaces/{ws_id}/rotas/{rota_id}/shifts/{shift_id}/copy",
        json={"target_date": "2026-11-03"},
        headers=headers,
    )
    assert copy_shift_resp.status_code == status.HTTP_400_BAD_REQUEST

    # 3i. Generate rota
    gen_resp = client.post(
        f"/api/v1/workspaces/{ws_id}/rotas/{rota_id}/generate",
        json={"time_limit_seconds": 5.0},
        headers=headers,
    )
    assert gen_resp.status_code == status.HTTP_400_BAD_REQUEST

    # 3j. Delete published rota directly
    del_rota_resp = client.delete(
        f"/api/v1/workspaces/{ws_id}/rotas/{rota_id}",
        headers=headers,
    )
    assert del_rota_resp.status_code == status.HTTP_400_BAD_REQUEST

    # 4. Unpublish rota back to draft
    unpub_resp = client.post(
        f"/api/v1/workspaces/{ws_id}/rotas/{rota_id}/unpublish",
        headers=headers,
    )
    assert unpub_resp.status_code == status.HTTP_200_OK
    assert unpub_resp.json()["status"] == "draft"
    assert unpub_resp.json()["published_at"] is None

    # Now edits are allowed again
    rename_ok = client.patch(
        f"/api/v1/workspaces/{ws_id}/rotas/{rota_id}",
        json={"name": "Renamed Draft Rota"},
        headers=headers,
    )
    assert rename_ok.status_code == status.HTTP_200_OK
    assert rename_ok.json()["name"] == "Renamed Draft Rota"


def test_archive_unarchive_and_confirmed_deletion(client: TestClient):
    """Test archiving retains records and unarchive clears publication metadata while preserving shift data."""
    headers, ws_id = _setup_user_and_workspace(client, "archive_test@example.com", "Archiver")

    staff_resp = client.post(
        f"/api/v1/workspaces/{ws_id}/staff",
        json={"first_name": "Carol", "last_name": "Danvers", "category": "full_time"},
        headers=headers,
    )
    staff_id = staff_resp.json()["id"]

    rota_resp = client.post(
        f"/api/v1/workspaces/{ws_id}/rotas",
        json={"name": "Archive Test Rota", "start_date": "2026-11-02", "end_date": "2026-11-08"},
        headers=headers,
    )
    rota_id = rota_resp.json()["id"]

    # Add a shift
    client.post(
        f"/api/v1/workspaces/{ws_id}/rotas/{rota_id}/shifts",
        json={
            "staff_id": staff_id,
            "date": "2026-11-02",
            "start_time": "09:00",
            "end_time": "17:00",
        },
        headers=headers,
    )

    # Publish rota
    approve_resp = client.post(
        f"/api/v1/workspaces/{ws_id}/rotas/{rota_id}/approve",
        headers=headers,
    )
    assert approve_resp.status_code == status.HTTP_200_OK
    assert approve_resp.json()["status"] == "published"
    assert approve_resp.json()["published_at"] is not None
    assert approve_resp.json()["published_by_user_id"] is not None

    # Archive published rota
    arch_resp = client.post(
        f"/api/v1/workspaces/{ws_id}/rotas/{rota_id}/archive",
        headers=headers,
    )
    assert arch_resp.status_code == status.HTTP_200_OK
    assert arch_resp.json()["status"] == "archived"
    assert arch_resp.json()["archived_at"] is not None

    # Unarchive back to draft - must clear archived_at AND published_at / published_by_user_id
    unarch_resp = client.post(
        f"/api/v1/workspaces/{ws_id}/rotas/{rota_id}/unarchive",
        headers=headers,
    )
    assert unarch_resp.status_code == status.HTTP_200_OK
    unarch_data = unarch_resp.json()
    assert unarch_data["status"] == "draft"
    assert unarch_data["archived_at"] is None
    assert unarch_data["published_at"] is None
    assert unarch_data["published_by_user_id"] is None
    # Historical shifts and child data must be retained intact
    assert len(unarch_data["shifts"]) == 1
    assert unarch_data["shifts"][0]["staff_id"] == staff_id

    # Delete draft rota
    del_resp = client.delete(
        f"/api/v1/workspaces/{ws_id}/rotas/{rota_id}",
        headers=headers,
    )
    assert del_resp.status_code == status.HTTP_204_NO_CONTENT

    # Confirm it is gone
    get_404 = client.get(
        f"/api/v1/workspaces/{ws_id}/rotas/{rota_id}",
        headers=headers,
    )
    assert get_404.status_code == status.HTTP_404_NOT_FOUND


def test_rota_duplication_workflow(client: TestClient):
    """Test duplication maps day offsets, handles collisions, preserves shifts/requirements, and omits audits/runs."""
    headers, ws_id = _setup_user_and_workspace(client, "duplicate_test@example.com", "Duplicator")

    staff_resp = client.post(
        f"/api/v1/workspaces/{ws_id}/staff",
        json={"first_name": "Diana", "last_name": "Prince", "category": "full_time"},
        headers=headers,
    )
    staff_id = staff_resp.json()["id"]

    # Source: 7 days Mon-Sun (2026-11-02 to 2026-11-08)
    rota_resp = client.post(
        f"/api/v1/workspaces/{ws_id}/rotas",
        json={"name": "Source Week", "start_date": "2026-11-02", "end_date": "2026-11-08"},
        headers=headers,
    )
    rota_id = rota_resp.json()["id"]

    # Add staffing requirement on Tuesday (2026-11-03, offset = 1)
    client.put(
        f"/api/v1/workspaces/{ws_id}/rotas/{rota_id}/requirements",
        json={
            "requirements": [
                {"date": "2026-11-03", "bands": [{"start_time": "09:00", "end_time": "17:00", "target_staff_count": 2}]}
            ]
        },
        headers=headers,
    )

    # Add overnight shift on Wednesday (2026-11-04, offset = 2, crosses midnight)
    client.post(
        f"/api/v1/workspaces/{ws_id}/rotas/{rota_id}/shifts",
        json={
            "staff_id": staff_id,
            "date": "2026-11-04",
            "start_time": "22:00",
            "end_time": "06:00",
            "notes": "Night shift",
        },
        headers=headers,
    )

    # 1. Target range duration mismatch (e.g. 5 days vs 7 days) -> HTTP 422
    mismatch_resp = client.post(
        f"/api/v1/workspaces/{ws_id}/rotas/{rota_id}/duplicate",
        json={
            "name": "Mismatched Week",
            "start_date": "2026-11-09",
            "end_date": "2026-11-13",  # only 5 days
        },
        headers=headers,
    )
    assert mismatch_resp.status_code == status.HTTP_422_UNPROCESSABLE_ENTITY

    # 2. Duplicate to next week (2026-11-09 to 2026-11-15, 7 days)
    dup_resp = client.post(
        f"/api/v1/workspaces/{ws_id}/rotas/{rota_id}/duplicate",
        json={
            "name": "Next Week Duplicated",
            "start_date": "2026-11-09",
            "end_date": "2026-11-15",
            "copy_requirements": True,
            "copy_shifts": True,
        },
        headers=headers,
    )
    assert dup_resp.status_code == status.HTTP_201_CREATED
    dup_data = dup_resp.json()
    assert dup_data["name"] == "Next Week Duplicated"
    assert dup_data["status"] == "draft"

    # Verify requirement mapped to Tuesday 2026-11-10 (offset 1)
    reqs = dup_data["requirements"]
    assert any(r["date"] == "2026-11-10" and len(r["bands"]) == 1 for r in reqs)

    # Verify shift mapped to Wednesday 2026-11-11 (offset 2)
    shifts = dup_data["shifts"]
    assert len(shifts) == 1
    assert shifts[0]["date"] == "2026-11-11"
    assert shifts[0]["start_time"] == "22:00"
    assert shifts[0]["end_time"] == "06:00"
    assert shifts[0]["notes"] == "Night shift"
    assert shifts[0]["active_override"] is None  # Audits must NOT be copied

    # 3. Collision check: Attempting to duplicate into the exact same active period -> HTTP 409
    collision_resp = client.post(
        f"/api/v1/workspaces/{ws_id}/rotas/{rota_id}/duplicate",
        json={
            "name": "Collision Attempt",
            "start_date": "2026-11-09",
            "end_date": "2026-11-15",
        },
        headers=headers,
    )
    assert collision_resp.status_code == status.HTTP_409_CONFLICT
    assert "already exists" in collision_resp.json()["detail"]


def test_cross_workspace_isolation_on_exports_and_lifecycle(client: TestClient):
    """Prove cross-workspace IDOR protection on exports, lifecycle, and duplication."""
    headers_a, ws_a = _setup_user_and_workspace(client, "user_a@example.com", "UserA")
    headers_b, ws_b = _setup_user_and_workspace(client, "user_b@example.com", "UserB")

    rota_a = client.post(
        f"/api/v1/workspaces/{ws_a}/rotas",
        json={"name": "Workspace A Rota", "start_date": "2026-11-02", "end_date": "2026-11-08"},
        headers=headers_a,
    ).json()
    rota_a_id = rota_a["id"]

    # User B attempts to export Rota A -> 403 or 404
    export_idor = client.get(
        f"/api/v1/workspaces/{ws_a}/rotas/{rota_a_id}/export?format=csv",
        headers=headers_b,
    )
    assert export_idor.status_code in (status.HTTP_403_FORBIDDEN, status.HTTP_404_NOT_FOUND)

    # User B attempts to approve Rota A -> 403 or 404
    approve_idor = client.post(
        f"/api/v1/workspaces/{ws_a}/rotas/{rota_a_id}/approve",
        headers=headers_b,
    )
    assert approve_idor.status_code in (status.HTTP_403_FORBIDDEN, status.HTTP_404_NOT_FOUND)

    # User B attempts to duplicate Rota A under Workspace B -> 404
    duplicate_idor = client.post(
        f"/api/v1/workspaces/{ws_b}/rotas/{rota_a_id}/duplicate",
        json={"name": "Stolen Rota", "start_date": "2026-11-09", "end_date": "2026-11-15"},
        headers=headers_b,
    )
    assert duplicate_idor.status_code == status.HTTP_404_NOT_FOUND
